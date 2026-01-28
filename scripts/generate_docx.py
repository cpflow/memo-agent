#!/usr/bin/env python3
"""Generate Word document from markdown."""
import sys
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def markdown_to_docx(md_path: str, output_path: str, template_path: str = None):
    """Convert markdown file to Word document.

    Args:
        md_path: Path to markdown file
        output_path: Path for output .docx file
        template_path: Optional path to template .docx for styling
    """
    doc = Document(template_path) if template_path else Document()

    with open(md_path) as f:
        content = f.read()

    # Remove YAML frontmatter if present
    content = re.sub(r"^---\n.*?\n---\n", "", content, flags=re.DOTALL)

    # Extract title (# heading)
    title_match = re.match(r"^# (.+)\n", content)
    if title_match:
        doc.add_heading(title_match.group(1), level=0)
        content = content[title_match.end() :]

    # Split by ## headers
    sections = re.split(r"^## ", content, flags=re.MULTILINE)

    for section in sections:
        section = section.strip()
        if not section:
            continue

        lines = section.split("\n")
        title = lines[0].strip()
        body = "\n".join(lines[1:]).strip()

        if title:
            doc.add_heading(title, level=1)

        # Process body paragraphs
        paragraphs = body.split("\n\n")
        for para in paragraphs:
            para = para.strip()
            if not para:
                continue

            # Handle bullet lists
            if para.startswith("- ") or para.startswith("* "):
                for item in para.split("\n"):
                    item = item.lstrip("-* ").strip()
                    if item:
                        doc.add_paragraph(item, style="List Bullet")

            # Handle numbered lists
            elif re.match(r"^\d+\. ", para):
                for item in para.split("\n"):
                    item = re.sub(r"^\d+\. ", "", item).strip()
                    if item:
                        doc.add_paragraph(item, style="List Number")

            # Handle ### subheadings
            elif para.startswith("### "):
                doc.add_heading(para[4:].strip(), level=2)

            # Regular paragraph
            else:
                # Clean up markdown formatting
                text = para.replace("\n", " ")
                text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)  # Bold
                text = re.sub(r"\*(.+?)\*", r"\1", text)  # Italic
                text = re.sub(r"\[Source: (.+?)\]", r"[\1]", text)  # Citations
                doc.add_paragraph(text)

    # Ensure output directory exists
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)

    doc.save(output_path)
    print(f"Saved to {output_path}")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python generate_docx.py <input.md> <output.docx> [template.docx]")
        sys.exit(1)

    md_path = sys.argv[1]
    output_path = sys.argv[2]
    template_path = sys.argv[3] if len(sys.argv) > 3 else None

    markdown_to_docx(md_path, output_path, template_path)

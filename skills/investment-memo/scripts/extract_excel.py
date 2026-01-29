#!/usr/bin/env python3
"""Extract data from Excel files using openpyxl."""
import sys
import io
from openpyxl import load_workbook

# Fix Windows encoding issues
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')


def extract(path: str, max_rows: int = 50) -> str:
    """Extract data from an Excel file as markdown tables.

    Args:
        path: Path to the Excel file
        max_rows: Maximum rows per sheet to extract (default 50)

    Returns:
        Extracted content as markdown tables
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    output_parts = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        output_parts.append(f"## Sheet: {sheet_name}\n")

        rows = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= max_rows:
                output_parts.append(f"\n[Truncated: showing {max_rows} rows]\n")
                break

            # Skip completely empty rows
            if all(cell is None for cell in row):
                continue

            # Convert cells to strings, handling None
            cells = [str(cell) if cell is not None else "" for cell in row]
            rows.append(cells)

        if rows:
            # Create markdown table
            # Use first row as header
            header = rows[0]
            output_parts.append("| " + " | ".join(header) + " |")
            output_parts.append("| " + " | ".join(["---"] * len(header)) + " |")

            for row in rows[1:]:
                # Pad row to match header length
                while len(row) < len(header):
                    row.append("")
                output_parts.append("| " + " | ".join(row[: len(header)]) + " |")

            output_parts.append("")  # Empty line between sheets

    wb.close()
    return "\n".join(output_parts)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python extract_excel.py <path-to-excel> [max-rows]")
        sys.exit(1)

    excel_path = sys.argv[1]
    max_rows = int(sys.argv[2]) if len(sys.argv) > 2 else 100
    print(extract(excel_path, max_rows))
